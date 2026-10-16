"""Generate DevGuard 10-slide presentation deck using python-pptx."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="DevGuard_Presentation.pptx"):
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # Blank

    # Color Palette (Dark Theme / Cyber-Defense Slate)
    BG_COLOR = RGBColor(15, 23, 42)       # #0F172A
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B
    BORDER_COLOR = RGBColor(51, 65, 85)   # #334155
    TEXT_MAIN = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    ACCENT_CYAN = RGBColor(56, 189, 248)  # #38BDF8
    ACCENT_INDIGO = RGBColor(129, 140, 248) # #818CF8
    ACCENT_GREEN = RGBColor(52, 211, 153) # #34D399
    ACCENT_AMBER = RGBColor(251, 191, 36) # #FBBF24

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, title, category=None):
        if category:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
            tf_cat = cat_box.text_frame
            tf_cat.word_wrap = True
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category.upper()
            p_cat.font.size = Pt(11)
            p_cat.font.bold = True
            p_cat.font.color.rgb = ACCENT_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7 if category else 0.5), Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

    def add_card(slide, left, top, width, height, title, items, badge_color=ACCENT_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        tx_box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = badge_color
        p_title.space_after = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "DEVGUARD"
    p0.font.size = Pt(56)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.space_after = Pt(8)

    p1 = tf1.add_paragraph()
    p1.text = "Local-First Privacy Guard & Autonomous Backend Repair Agent"
    p1.font.size = Pt(24)
    p1.font.color.rgb = TEXT_MAIN
    p1.space_after = Pt(16)

    p2 = tf1.add_paragraph()
    p2.text = "Powered entirely on-device by Gemma 4 E4B (Ollama) | Zero Screen/Key Recording | Fail-Closed Security"
    p2.font.size = Pt(14)
    p2.font.color.rgb = ACCENT_INDIGO

    # Info footer
    foot = s1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.333), Inches(0.8))
    p_foot = foot.text_frame.paragraphs[0]
    p_foot.text = "Hackathon Presentation | Track: Best Use of Gemma 4 & Open Source AI | 2026"
    p_foot.font.size = Pt(12)
    p_foot.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: The Core Problem
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "The Critical Security Blindspot in AI Coding Assistants", "Problem Statement")

    add_card(s2, 0.8, 1.8, 3.7, 4.8, "1. Accidental Cloud Leaks", [
        "Developers inadvertently paste proprietary API keys, database URLs, and AWS tokens into AI chats.",
        "Diagnostic stack traces contain real customer PII and internal microservice URIs.",
        "Cloud models retain prompts, risking intellectual property and compliance violations."
    ], ACCENT_AMBER)

    add_card(s2, 4.8, 1.8, 3.7, 4.8, "2. Invasive Surveillance", [
        "Traditional enterprise security relies on screen scrapers, video watchers, or keyboard loggers.",
        "Heavy privacy intrusion for developers and massive CPU/battery performance drain.",
        "Lacks semantic understanding of developer workflows and code facts."
    ], ACCENT_AMBER)

    add_card(s2, 8.8, 1.8, 3.7, 4.8, "3. Slow & Risky Bug Fixing", [
        "Debugging production regressions requires manual log copy-pasting to cloud assistants.",
        "Manual patch application often overwrites concurrent developer edits (lost changes).",
        "Untested AI patches can introduce subtle regression failures."
    ], ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 3: DevGuard Dual-Agent Solution & Architecture
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "DevGuard Architecture: Dual-Agent Local-First Protection", "Solution Overview")

    add_card(s3, 0.8, 1.8, 5.6, 4.8, "Agent 1: Synchronous Privacy Gate", [
        "Interceps text prompts & clipboard BEFORE external delivery.",
        "High-priority foreground queue powered by local Gemma 4 E4B.",
        "High-entropy Gitleaks pattern detection + semantic confidential prose analysis.",
        "Exact-span type-safe redaction ([CUSTOMER_EMAIL_1], [REDACTED_SECRET_TOKEN_1]).",
        "Fail-Closed Gate: If offline, holds delivery (REVIEW) rather than leaking data."
    ], ACCENT_CYAN)

    add_card(s3, 6.8, 1.8, 5.6, 4.8, "Agent 2: Autonomous Sandbox Repair", [
        "Investigates backend test failures in disposable Docker sandboxes.",
        "Zero screen/audio/video recording: uses structured pytest logs & code diffs only.",
        "Gemma 4 explores code with scoped tools & produces minimal unified diffs.",
        "Verified against independent, immutable test oracles.",
        "Base-hash conflict protection: Rejects patch if working tree changed (HTTP 409)."
    ], ACCENT_INDIGO)

    # -------------------------------------------------------------
    # SLIDE 4: Feature 1 - Precision Pre-Delivery Privacy Gate
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "Feature 1: Exact-Span Redaction & Fail-Closed Gate", "Core Feature")

    add_card(s4, 0.8, 1.8, 5.6, 4.8, "Semantic & Pattern Redaction", [
        "Multi-layer detection: Regex scans for known token signatures + Gemma 4 contextual reasoning.",
        "Catches confidential text that regex misses: unformatted passwords, internal hostnames, and PII.",
        "Type-Safe Placeholders: Replaces secrets while keeping 100% of the debugging structure.",
        "Preserves line numbers, exception classes, and syntax so AI can still fix the bug."
    ], ACCENT_CYAN)

    add_card(s4, 6.8, 1.8, 5.6, 4.8, "Deterministic Gate Policy", [
        "ALLOW: Clean text is approved for transmission with verifiable receipt.",
        "SANITIZE: Exact spans replaced with safe placeholders before release.",
        "REVIEW: Triggered on model offline or ambiguities; prevents uninspected leakage.",
        "BLOCK: Reject malicious inputs (directory traversal, command injection).",
        "One-Time Token Receipts: Delivered payloads cannot be re-transmitted or intercepted."
    ], ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 5: Feature 2 - Opt-In macOS Clipboard Watcher
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "Feature 2: Native Text Clipboard Protection & Cleaned Copy", "Core Feature")

    add_card(s5, 0.8, 1.8, 5.6, 4.8, "Strict Opt-In & Privacy Invariants", [
        "Explicit Consent: Capture starts strictly OFF; requires local user activation.",
        "Text Payloads Only: Ignores images, binaries, and unmanaged clipboard buffers.",
        "Anti-Scan Loop: Hash tracking prevents infinite re-scans when copying sanitized text.",
        "Transient Processing: Clipboard contents are never written to disk or audit logs.",
        "Rapid Overwrite Cancel: If user copies again, previous inspection aborts cleanly."
    ], ACCENT_CYAN)

    add_card(s5, 6.8, 1.8, 5.6, 4.8, "One-Click 'Copy Cleaned Text'", [
        "Detects connection strings (DATABASE_URL='postgres://admin:pass@host:5432').",
        "Instantly displays sanitized safe preview in operator dashboard.",
        "Operator clicks 'Copy Cleaned Text' to replace sensitive clipboard contents.",
        "Developer pastes directly into AI assistant without leaking database credentials.",
        "Audited Metadata: Only event hashes and redaction counts survive restarts."
    ], ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 6: Feature 3 - Autonomous Docker Sandbox Repair
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Feature 3: Background Code Repair via Sandboxed Oracles", "Core Feature")

    add_card(s6, 0.8, 1.8, 5.6, 4.8, "Reproducible Sandbox Diagnosis", [
        "Triggered by test failure signals (e.g., cart quantity calculation regression).",
        "Executes inside isolated, pinned Docker containers with immutable tests.",
        "Gemma 4 uses strictly scoped tools: 'read_allowed_file', 'propose_patch', 'run_tests'.",
        "Zero Host Access: Cannot touch host filesystem or run uncontained shell scripts.",
        "Step Budget: Enforces strict 6-step limits to prevent infinite agent loops."
    ], ACCENT_INDIGO)

    add_card(s6, 6.8, 1.8, 5.6, 4.8, "Independent Test Verification", [
        "Baseline Check: Verifies regression genuinely reproduces in clean container.",
        "Independent Oracle: Evaluates candidate patch against protected 'trusted_tests'.",
        "No Fabricated Outcomes: Fails closed if test suite fails or error persists.",
        "Generates clean unified diff ready for human operator review.",
        "Complete transparency: Model actions and verifier outputs clearly logged."
    ], ACCENT_CYAN)

    # -------------------------------------------------------------
    # SLIDE 7: Feature 4 - Concurrency Safety & Stale Base-Hash Guard
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Feature 4: Concurrency Protection & Stale Patch Rejection", "Core Feature")

    add_card(s7, 0.8, 1.8, 5.6, 4.8, "Cryptographic Base-Hash Validation", [
        "When an incident is captured, SHA-256 hashes of target files are recorded.",
        "Repair agent builds patch against that specific snapshot revision.",
        "Before patch application, DevGuard re-hashes the live working tree on disk.",
        "If developer edited files while the repair ran -> Base hashes mismatch!",
        "Result: Instant HTTP 409 Conflict ('STALE_CONTENT') — zero code overwrites."
    ], ACCENT_AMBER)

    add_card(s7, 6.8, 1.8, 5.6, 4.8, "Strict Operator Authority", [
        "AI Proposes, Human Authorizes: No automated silent code modifications.",
        "Idempotent Application: Retrying an already applied patch returns cached state.",
        "File-Level Locking: Cross-thread file locks prevent race conditions during apply.",
        "Clean Reversion: Unapplied or canceled candidates leave original files untouched.",
        "Full Git compatibility: Operates cleanly on standard git branches and diffs."
    ], ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 8: Operator Dashboard & Real-Time Observability
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Operator Dashboard: Complete Transparency & Zero Fluff", "User Experience")

    add_card(s8, 0.8, 1.8, 3.7, 4.8, "5 Unified Views", [
        "🛡️ Privacy Gate: Real-time prompt evaluation with scenario prefill selectors.",
        "📋 Clipboard View: Live capture status, safe previews & cleaned copy button.",
        "🔍 Coverage Matrix: Honest VERIFIED / PARTIAL / UNSUPPORTED boundaries.",
        "🔧 Patch Review: Interactive side-by-side unified diffs & verification status.",
        "📜 Audit Stream: Filtered chronological metadata events (no raw secrets)."
    ], ACCENT_CYAN)

    add_card(s8, 4.8, 1.8, 3.7, 4.8, "Honest Fixture Mode", [
        "Prominent warning banner when running in simulated fixture mode.",
        "Declared synthetic data never masquerades as live model inference.",
        "Clearly distinguishes offline demonstration from live on-device evaluation.",
        "Full failure procedure documented for judges and evaluators."
    ], ACCENT_AMBER)

    add_card(s8, 8.8, 1.8, 3.7, 4.8, "Enterprise Compliance", [
        "Zero Token Leaks: Operator session tokens kept server-side, never in UI states.",
        "Cross-Platform: Windows operator viewer + native macOS capture host.",
        "FastAPI Backend: Strictly typed Pydantic models with 32 KiB bounded requests.",
        "Durable SQLite Audit: Up to 5,000 events survive restart without private data."
    ], ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 9: Real-World Uses & Target Use Cases
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Target Use Cases: Where DevGuard Delivers Immediate Value", "Use Cases & Impact")

    add_card(s9, 0.8, 1.8, 3.7, 4.8, "Enterprise AI Adoption", [
        "Target: Engineering teams adopting Cursor, Copilot, or Claude Code.",
        "Value: Prevent accidental leakage of AWS keys, database strings, and internal hostnames.",
        "Compliance: Adhere to SOC 2, ISO 27001, and enterprise security policies seamlessly."
    ], ACCENT_CYAN)

    add_card(s9, 4.8, 1.8, 3.7, 4.8, "FinTech & Healthcare", [
        "Target: Developers working with HIPAA PII or PCI-DSS financial transactions.",
        "Value: Automatic exact-span redaction of customer emails, account IDs, and patient records in crash logs.",
        "Context Retained: AI can fix the crash without seeing regulated health or payment data."
    ], ACCENT_INDIGO)

    add_card(s9, 8.8, 1.8, 3.7, 4.8, "Autonomous CI/CD Triage", [
        "Target: Fast-moving agile teams managing regression test failures.",
        "Value: Automatically spins up isolated Docker sandboxes when tests break, diagnoses root cause, and prepares verified diff.",
        "Time Saved: Reduces bug resolution from hours of context-switching to a 1-click review."
    ], ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 10: Tech Stack & Key Takeaways
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)
    add_header(s10, "Technology Stack & Why DevGuard Wins", "Summary")

    add_card(s10, 0.8, 1.8, 5.6, 4.8, "The Modern Local Tech Stack", [
        "Local LLM: Gemma 4 E4B (q4_K_M) running via Ollama — 100% on-device inference.",
        "Backend Engine: FastAPI + Pydantic v2 + Python 3.11+ async loop.",
        "Security Scanner: Gitleaks integration + custom semantic span matcher.",
        "Verification Sandbox: Docker containerized runner with pinned test images.",
        "Operator UI: Streamlit 1.40+ dashboard with reactive HTTP display client.",
        "Storage: File-locked SQLite metadata store with POSIX permission guards."
    ], ACCENT_INDIGO)

    add_card(s10, 6.8, 1.8, 5.6, 4.8, "Core Invariants (Why DevGuard Wins)", [
        "1. True Dual-Agent Synergy: Fast synchronous gate + background sandbox healer.",
        "2. Zero Surveillance: No screens, keyloggers, or audio scraping — only code facts.",
        "3. Fail-Closed Privacy: Offline models hold requests; no data escapes uninspected.",
        "4. Cryptographic Concurrency Safety: Stale base-hashes reject conflicting writes.",
        "5. Complete Honesty: Clear capability matrix, no fabricated AI outcomes."
    ], ACCENT_CYAN)

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
